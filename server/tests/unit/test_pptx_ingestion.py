"""Native PowerPoint tables must survive parsing with their original column positions."""
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")
from pptx.util import Inches  # noqa: E402

from yantra_server.knowledge.ingest.chunk import chunk_document  # noqa: E402
from yantra_server.knowledge.ingest.parse import parse_document  # noqa: E402


def test_table_only_slide_preserves_identifiers_blanks_and_numbers(tmp_path: Path) -> None:
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(3, 4, Inches(1), Inches(1), Inches(8), Inches(2)).table
    values = [
        ["Vendor", "Price INR", "Missing", "Material"],
        ["SYN-A", "480000", "", "SiC/SiC"],
        ["SYN-B", "445000", "UNKNOWN", "Carbon/SiC"],
    ]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    path = tmp_path / "native.pptx"
    deck.save(path)
    parsed = parse_document(path)
    block = parsed.all_blocks()[0]
    assert block.kind == "table" and block.page == 1
    assert block.meta["rows"] == values
    assert "SYN-A | 480000 |  | SiC/SiC" in block.text
    chunks = [c for c in chunk_document(parsed) if not c.meta.get("is_parent")]
    assert any("SYN-B | 445000 | UNKNOWN | Carbon/SiC" in c.text for c in chunks)
    assert all(c.page_start == 1 for c in chunks)


def test_nested_group_and_merged_header_have_slide_locator(tmp_path: Path) -> None:
    deck = pptx.Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Review findings"
    shape = slide.shapes.add_table(3, 3, 0, Inches(1), Inches(8), Inches(2))
    table = shape.table
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "SYNTHETIC INSPECTION"
    table.cell(1, 0).text = "W-101"
    table.cell(1, 1).text = "accepted"
    table.cell(2, 0).text = "W-I01"
    table.cell(2, 1).text = "not issued"
    # PowerPoint allows a table in a group, although add_table lives on SlideShapes.
    nested.shapes._spTree.append(shape._element)
    path = tmp_path / "grouped.pptx"
    deck.save(path)
    parsed = parse_document(path)
    assert not parsed.pages[0].blocks
    assert "Review findings" in parsed.full_text()
    block = next(b for b in parsed.all_blocks() if b.kind == "table")
    assert block.page == 2 and block.section_path.startswith("Slide 2 / Group")
    assert block.text.count("SYNTHETIC INSPECTION") == 1
    assert block.meta["merged_cells"] == [
        {"row": 1, "column": 1, "row_span": 1, "column_span": 3}
    ]
    assert "W-101 | accepted" in block.text
    assert "W-I01 | not issued" in block.text
