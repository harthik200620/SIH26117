from types import SimpleNamespace

import openpyxl
import pytest

from yantra_server.render.service import RenderError, RenderService
from yantra_server.render.workbook_checks import check_workbook
from yantra_server.tools.base import ToolContext
from yantra_server.tools.builtin.render import RenderDocumentArgs, RenderDocumentTool


@pytest.mark.parametrize(
    "formula",
    [
        "=basic_INR+freight_INR",
        "=#REF!+1",
        "=Missing!A1",
        "='[outside.xlsx]Sheet1'!A1",
        '=WEBSERVICE("https://example.com")',
        "=A2",
        "=XFE1",
        "=SUM(A1:A1048576)",
    ],
)
def test_broken_or_external_formulas_fail(formula):
    book = openpyxl.Workbook()
    book.active["A2"] = formula
    with pytest.raises(ValueError):
        check_workbook(book)


def test_valid_cross_sheet_arithmetic_and_indirect_cycle():
    book = openpyxl.Workbook()
    book.active.title = "Input data"
    book.active.append([10, 2])
    ws = book.create_sheet("Totals")
    ws["A1"] = "=SUM('Input data'!$A$1:$B$1)"
    ws["A2"] = "=A1*(1+0.18)"
    assert check_workbook(book) == 2
    book.active["A1"] = "=Totals!A2"
    with pytest.raises(ValueError, match="Circular"):
        check_workbook(book)


def test_distant_blank_references_do_not_expand_workbook():
    book = openpyxl.Workbook()
    book.active["A1"] = "=XFD1048576"
    assert check_workbook(book) == 1
    assert book.active.max_row == 1 and book.active.max_column == 1


def test_undefined_column_feedback_identifies_real_cell_addresses():
    book = openpyxl.Workbook()
    book.active.append(["Basic", "Freight", "Total"])
    book.active.append([10, 2, "=Basic+Freight"])
    with pytest.raises(ValueError, match="Basic=A2, Freight=B2, Total=C2"):
        check_workbook(book)


def test_rejected_render_preserves_existing_output_and_provenance(tmp_path):
    service = RenderService(tmp_path)
    path = tmp_path / "report.xlsx"
    payload = {"sheets": [{"columns": ["Price", "Total"], "rows": [[10, "=A2*1.18"]]}]}
    service.render(doc_type="xlsx", schema_id="data_table", out_path=path, data=payload)
    before = path.read_bytes()
    provenance = path.with_suffix(".xlsx.provenance.json").read_bytes()
    payload["sheets"][0]["rows"][0][1] = "=price*1.18"
    with pytest.raises(ValueError, match="undefined"):
        service.render(doc_type="xlsx", schema_id="data_table", out_path=path, data=payload)
    assert path.read_bytes() == before
    assert path.with_suffix(".xlsx.provenance.json").read_bytes() == provenance


def test_tables_never_silently_drop_cells_and_evidence_is_preserved(tmp_path):
    service = RenderService(tmp_path)
    with pytest.raises(RenderError, match="Row 1"):
        service.validate("data_table", {"sheets": [{"columns": ["A"], "rows": [[1, 2]]}]})
    path = tmp_path / "offers.xlsx"
    service.render(
        doc_type="xlsx",
        schema_id="data_table",
        out_path=path,
        data={
            "sheets": [
                {
                    "caption": "Offers / inputs",
                    "columns": ["Equipment ID", "Amount (INR)"],
                    "rows": [["001", 100]],
                }
            ],
            "summary": "Approval pending",
            "appendix": {"citations": ["input.csv row 2"], "assumptions": ["Synthetic fixture"]},
        },
    )
    book = openpyxl.load_workbook(path)
    assert book.active["A2"].value == "001"
    assert book.active["B2"].data_type == "n"
    assert book.active["B1"].alignment.wrap_text
    assert book.active.freeze_panes == "A2"
    values = [c.value for row in book["Decision and evidence"] for c in row]
    assert {"Approval pending", "input.csv row 2", "Synthetic fixture"}.issubset(values)


@pytest.mark.parametrize(
    "schema,payload",
    [
        (
            "report",
            {
                "title": "Report",
                "sections": [{"heading": "Figure", "figures": [{"path": "../private.png"}]}],
            },
        ),
        (
            "presentation",
            {"title": "Deck", "slides": [{"title": "Figure", "figure": "../private.png"}]},
        ),
        ("pid_review", {"title": "Review", "overlay_figure": "../private.png"}),
    ],
)
async def test_embedded_images_cannot_escape_workspace(tmp_path, monkeypatch, schema, payload):
    monkeypatch.setattr(
        "yantra_server.tools.builtin.render._render_service", lambda ctx: RenderService(tmp_path)
    )
    ctx = ToolContext(
        workspace=tmp_path,
        state=SimpleNamespace(audit=SimpleNamespace(head=lambda: None)),
        sandbox=None,
    )
    result = await RenderDocumentTool().run(
        RenderDocumentArgs(
            type="docx",
            schema_id=schema,
            data_json=payload,
            out_path="out.docx",
        ),
        ctx,
    )
    assert not result.ok and "escapes the workspace" in result.error
    assert not (tmp_path / "out.docx").exists()


def test_json_file_must_be_an_object(tmp_path):
    (tmp_path / "list.json").write_text("[]")
    ctx = SimpleNamespace(resolve_path=lambda name: tmp_path / name)
    assert RenderDocumentTool()._load_data("list.json", ctx) is None
