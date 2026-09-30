"""Static workbook checks. These do not calculate formulas or certify their meaning."""

from __future__ import annotations

import re
from typing import Any

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import get_column_letter, range_boundaries

MAX_CELLS = 200_000
MAX_RANGE_CELLS = 20_000
EXTERNAL_FUNCTIONS = {"WEBSERVICE", "HYPERLINK", "RTD", "DDE", "CALL", "REGISTER.ID"}


def check_workbook(book: Any) -> int:
    """Reject broken references, errors, external links and cycles without execution."""
    if not book.sheetnames:
        raise ValueError("Workbook contains no worksheets")
    if getattr(book, "_external_links", []):
        raise ValueError("External workbook links are not allowed")
    if sum(s.max_row * s.max_column for s in book) > MAX_CELLS:
        raise ValueError(f"Workbook exceeds {MAX_CELLS} cells for static validation")
    sheets = {s.title.casefold(): s for s in book}
    formula_cells = {
        (s.title.casefold(), c.coordinate)
        for s in book
        for row in s
        for c in row
        if c.data_type == "f"
    }
    graph: dict[tuple[str, str], set[tuple[str, str]]] = {}
    edges = 0
    for sheet in book:
        names = {n.casefold() for n in set(book.defined_names) | set(sheet.defined_names)}
        for row in sheet:
            for cell in row:
                where = f"{sheet.title}!{cell.coordinate}"
                if cell.data_type == "e":
                    raise ValueError(f"{where}: Excel error {cell.value}")
                if cell.data_type != "f":
                    continue
                key = (sheet.title.casefold(), cell.coordinate)
                graph[key] = set()
                for token in Tokenizer(cell.value).items:
                    value = token.value
                    if token.type == "FUNC" and token.subtype == "OPEN":
                        function = value[:-1].upper().removeprefix("_XLFN.")
                        if function in EXTERNAL_FUNCTIONS:
                            raise ValueError(
                                f"{where}: external/active function {function} is not allowed"
                            )
                    if token.type == "OPERAND" and token.subtype == "ERROR":
                        raise ValueError(f"{where}: formula contains {value}")
                    if token.type != "OPERAND" or token.subtype != "RANGE":
                        continue
                    if "[" in value or "]" in value or "|" in value:
                        # Structured references require a separate validator; never
                        # mistake them for verified ordinary cell references.
                        raise ValueError(
                            f"{where}: external or structured reference requires review: {value}"
                        )
                    target = sheet
                    if "!" in value:
                        title, value = value.rsplit("!", 1)
                        title = title.strip("'").replace("''", "'")
                        if title.casefold() not in sheets:
                            raise ValueError(f"{where}: unknown worksheet {title}")
                        target = sheets[title.casefold()]
                    if not re.fullmatch(
                        r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?", value
                    ):
                        if value.casefold() in names:
                            raise ValueError(
                                f"{where}: named reference {value} needs calculation-engine validation"
                            )
                        references = ", ".join(
                            f"{str(header.value)[:60]}={header.column_letter}{cell.row}"
                            for header in sheet[1][:32]
                            if header.value is not None
                        )
                        raise ValueError(
                            f"{where}: undefined formula name {value} or unsupported reference. "
                            f"Column labels are not Excel variables. Cell references for this row: {references}. "
                            "Rewrite formulas using the applicable cell addresses."
                        )
                    left, top, right, bottom = range_boundaries(value)
                    if right > 16384 or bottom > 1048576 or left > right or top > bottom:
                        raise ValueError(f"{where}: invalid Excel cell range {value}")
                    size = (right - left + 1) * (bottom - top + 1)
                    edges += size
                    if size > MAX_RANGE_CELLS or edges > MAX_CELLS:
                        raise ValueError(f"{where}: formula dependencies exceed validation budget")
                    # Checking a distant blank reference must not materialize cells
                    # or expand the worksheet's used range to a million rows.
                    for row_number in range(top, bottom + 1):
                        for column in range(left, right + 1):
                            dependency = (
                                target.title.casefold(),
                                f"{get_column_letter(column)}{row_number}",
                            )
                            if dependency in formula_cells:
                                graph[key].add(dependency)
    # Iterative topological elimination avoids recursion limits on long sheets.
    reverse: dict[tuple[str, str], list[tuple[str, str]]] = {}
    counts = {cell: len(deps) for cell, deps in graph.items()}
    for cell, dependencies in graph.items():
        for dependency in dependencies:
            reverse.setdefault(dependency, []).append(cell)
    ready = [cell for cell, count in counts.items() if count == 0]
    visited = 0
    while ready:
        cell = ready.pop()
        visited += 1
        for dependent in reverse.get(cell, []):
            counts[dependent] -= 1
            if counts[dependent] == 0:
                ready.append(dependent)
    if visited != len(graph):
        cell = next(cell for cell, count in counts.items() if count)
        raise ValueError(f"Circular formula dependency involving {cell[0]}!{cell[1]}")
    return len(graph)
