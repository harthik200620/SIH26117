"""Read captured workflow files without modifying their bytes or assigning a model score."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "folder", type=Path, help="Captured run folder containing run.json and workspace/"
    )
    args = parser.parse_args()
    report = json.loads((args.folder / "run.json").read_text(encoding="utf-8"))
    artifacts = []
    for path in sorted((args.folder / "workspace").glob("*")):
        if path.suffix.lower() != ".xlsx":
            continue
        import openpyxl

        artifact = {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        try:
            book = openpyxl.load_workbook(path, data_only=False, read_only=True)
            try:
                artifact["sheets"] = [
                    {
                        "name": sheet.title,
                        "rows": [
                            [
                                {
                                    "address": cell.coordinate,
                                    "type": cell.data_type,
                                    "value": cell.value,
                                }
                                for cell in row
                                if cell.value is not None
                            ]
                            for row in sheet.iter_rows(
                                max_row=min(sheet.max_row, 200), max_col=min(sheet.max_column, 128)
                            )
                        ],
                        "truncated": sheet.max_row > 200 or sheet.max_column > 128,
                    }
                    for sheet in book
                ]
            finally:
                book.close()
        except Exception as exc:
            artifact["error"] = str(exc)
        artifacts.append(artifact)
    result = {
        "run_id": report["run_id"],
        "application_status": report["run"]["status"],
        "source_bytes_unchanged": report["source_bytes_unchanged"],
        "implementation_unchanged": report["implementation_unchanged"],
        "artifacts": artifacts,
        "scope": "Raw cell inspection. Formula text is not recalculated output; no semantic pass is assigned.",
    }
    output = args.folder / "artifact_inspection.json"
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"output": str(output), "artifacts": len(artifacts)}))


if __name__ == "__main__":
    main()
