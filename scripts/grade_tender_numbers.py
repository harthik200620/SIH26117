"""Independent numeric checks for the synthetic tender fixture, after engine recalculation.

Run outside the model workspace. This does not grade narrative decisions or approvals.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import openpyxl


def grade(folder: Path) -> dict:
    report = json.loads((folder / "run.json").read_text(encoding="utf-8"))
    source = folder / "workspace/offers.csv"
    workbook = folder / "workspace/tender-comparison.xlsx"
    recalculation = folder / "recalculation/recalculated-values.json"
    checks = []

    def check(name: str, passed: bool, detail: str) -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    check("terminal_done", report["run"]["status"] == "done", report["run"]["status"])
    check("sources_unchanged", report["source_bytes_unchanged"], "Compared with run input hashes")
    check(
        "implementation_unchanged",
        report["implementation_unchanged"],
        "Compared recorded solver file hashes",
    )
    try:
        engine_record = json.loads((folder / "recalculation/engine-record.json").read_text())
        check(
            "recalculation_source_hash",
            engine_record["inputSha256"] == hashlib.sha256(workbook.read_bytes()).hexdigest(),
            "Recalculated the captured original workbook bytes",
        )
        with source.open(encoding="utf-8-sig", newline="") as stream:
            offers = list(csv.DictReader(stream))
        values = json.loads(recalculation.read_text(encoding="utf-8"))
        book = openpyxl.load_workbook(workbook, data_only=False, read_only=True)
        try:
            sheet = book.worksheets[0]
            headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
            check(
                "recalculation_headers",
                values[0][: len(headers)] == headers,
                "Engine output aligns with original headers",
            )
            positions = {str(name): i for i, name in enumerate(headers)}
            missing = set(offers[0]) - positions.keys()
            check(
                "source_columns", not missing, f"Missing original source headers: {sorted(missing)}"
            )
            if missing:
                raise ValueError("Cannot align original source columns")
            totals = [
                i
                for i, label in enumerate(headers)
                if "landed" in str(label).lower() and "total" in str(label).lower()
            ]
            if len(totals) != 1:
                raise ValueError("Expected one uniquely labelled landed total column")
            total_column = totals[0]
            numeric = {"basic_INR", "freight_INR", "tax_fraction", "delivery_weeks"}
            original = {offer["vendor"]: offer for offer in offers}
            seen = []
            for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
                vendor = cells[positions["vendor"]].value
                if vendor not in original:
                    check(f"row_{row_number}_vendor", False, f"Unexpected vendor {vendor!r}")
                    continue
                seen.append(vendor)
                offer = original[vendor]
                for name, value in offer.items():
                    actual = cells[positions[name]]
                    if name in numeric:
                        same = actual.data_type == "n" and Decimal(str(actual.value)) == Decimal(
                            value
                        )
                    else:
                        same = actual.data_type == "s" and actual.value == value
                    check(f"{vendor}_{name}", same, f"Observed {actual.value!r}; source {value!r}")
                actual_total = values[row_number - 1][total_column]
                expected = (Decimal(offer["basic_INR"]) + Decimal(offer["freight_INR"])) * (
                    1 + Decimal(offer["tax_fraction"])
                )
                check(
                    f"{vendor}_landed_formula",
                    cells[total_column].data_type == "f",
                    str(cells[total_column].value),
                )
                check(
                    f"{vendor}_landed_value",
                    isinstance(actual_total, (int, float))
                    and abs(Decimal(str(actual_total)) - expected) <= Decimal("0.01"),
                    f"Engine {actual_total!r}; independently expected {expected} INR",
                )
            check(
                "offer_coverage", sorted(seen) == sorted(original), f"Observed vendor rows: {seen}"
            )
        finally:
            book.close()
    except Exception as exc:
        check("inspection_complete", False, str(exc))
    return {
        "run_id": report["run_id"],
        "numeric_checks_passed": all(c["passed"] for c in checks),
        "checks": checks,
        "sha256": {
            p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in (source, workbook, recalculation)
            if p.is_file()
        },
        "scope": "Known synthetic tender rule: (basic + freight) * (1 + tax fraction). Values from separate Artifact Tool recalculation. Narrative selection, exclusions, citations and approval still require inspection. Not an independent expert or unseen evaluation.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    args = parser.parse_args()
    result = grade(args.folder)
    (args.folder / "numeric-checks.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "run_id": result["run_id"],
                "numeric_checks_passed": result["numeric_checks_passed"],
                "checks": len(result["checks"]),
            }
        )
    )


if __name__ == "__main__":
    main()
